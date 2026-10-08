"""Parser dell'export Excel mensile di mySET (endpoint
/consumption/excel, vedi api.SetApiClient.async_get_consumption_excel_month).

Formato CONFERMATO su dati reali (30/09/2026, agosto 2026: 2976 righe =
31 giorni x 96 quarti d'ora, somma dei valori identica al millesimo alla
somma dei `total` della vista mese JSON):

  foglio "Energia Attiva Prevelata" (sic, refuso del portale):
    riga 1: POD | Ruolo | Unita' di misura | Anno | Mese
    riga 2: <POD> | Energia Attiva Prevelata | kWh | 2026 | agosto
    riga 3: vuota
    riga 4: Giorno | Da ora | A ora | Valore | Costante K | Matricola Contatore
    poi:    01/08/2026 | 00:00 | 00:15 | 0,140 | 1 | <matricola>   (x 96 al giorno)
  foglio "Energia Reattiva Q1": stesso layout, energia reattiva - ignorato.

Tutte le celle sono stringhe: date dd/mm/yyyy, ore HH:MM, valori con la
virgola decimale. Si converte ogni riga in un quarto d'ora e si
ricompongono elementi "ora" nella STESSA forma della vista giorno JSON
({year, month, day, hour, quarters[4], total, estimated, kConstant}),
cosi' l'import passa dall'unico percorso gia' testato
(statistics.async_import_curve).

openpyxl e' una dipendenza dichiarata nel manifest (a differenza di
pcf_common, dove l'import xlsx e' una utility mai usata dal coordinator).
Il parsing e' sincrono e su ~3000 righe non e' istantaneo: va eseguito
fuori dall'event loop (hass.async_add_executor_job).

Cambio ora legale: NON verificato. Il giorno da 25 ore avra' presumibilmente
8 righe con la stessa "Da ora" tra le 02:00 e le 03:00: qui finiscono
sommate nello stesso elemento ora (il totale del giorno resta giusto,
l'attribuzione oraria di quell'ora e' approssimata) - stesso limite della
vista giorno JSON, vedi statistics.py.
"""
from __future__ import annotations

import io
import logging
from collections import defaultdict
from datetime import date, datetime
from typing import Any

_LOGGER = logging.getLogger(__name__)

FOGLIO_ATTIVA = "Energia Attiva Prevelata"
_INTESTAZIONI_ATTESE = ("giorno", "da ora", "valore")


class SetExcelError(Exception):
    """File non parsabile o con un layout diverso da quello confermato."""


def _numero(valore: Any) -> float | None:
    if valore is None or valore == "":
        return None
    try:
        return float(str(valore).strip().replace(",", "."))
    except ValueError:
        return None


def parse_excel_mensile(xlsx: bytes) -> dict[date, list[dict[str, Any]]]:
    """xlsx (bytes) -> {giorno: [24 elementi nella forma della vista giorno]}.

    Solleva SetExcelError se manca il foglio dell'energia attiva o la riga
    di intestazione attesa. Righe singole non parsabili vengono scartate
    con un log, non fanno fallire il mese."""
    import openpyxl  # dipendenza dichiarata nel manifest, import qui per non pesare all'avvio

    try:
        wb = openpyxl.load_workbook(io.BytesIO(xlsx), data_only=True, read_only=True)
    except Exception as err:  # noqa: BLE001 - openpyxl solleva tipi eterogenei (zipfile, KeyError, ...)
        raise SetExcelError(f"File Excel non leggibile: {err}") from err

    foglio = next((ws for ws in wb.worksheets if ws.title.strip().lower() == FOGLIO_ATTIVA.lower()), None)
    if foglio is None:
        # Ripiego: il primo foglio con "attiva" nel nome - ma non "reattiva"
        # (il secondo foglio e' "Energia Reattiva Q1").
        foglio = next(
            (
                ws for ws in wb.worksheets
                if "attiva" in ws.title.lower() and "reattiva" not in ws.title.lower()
            ),
            None,
        )
    if foglio is None:
        raise SetExcelError(f"Foglio {FOGLIO_ATTIVA!r} non trovato (fogli: {wb.sheetnames})")

    righe = foglio.iter_rows(values_only=True)
    colonne: dict[str, int] | None = None
    for riga in righe:
        etichette = [str(c).strip().lower() if c is not None else "" for c in riga]
        if all(e in etichette for e in _INTESTAZIONI_ATTESE):
            colonne = {e: etichette.index(e) for e in etichette if e}
            break
    if colonne is None:
        raise SetExcelError("Riga di intestazione 'Giorno | Da ora | ... | Valore' non trovata")

    i_giorno, i_da, i_val = colonne["giorno"], colonne["da ora"], colonne["valore"]
    i_k = colonne.get("costante k")

    # (giorno, ora) -> [valori dei quarti], in ordine di apparizione
    quarti: dict[tuple[date, int], list[float]] = defaultdict(list)
    costanti: dict[date, str] = {}
    scartate = 0
    for riga in righe:
        if not riga or len(riga) <= max(i_giorno, i_da, i_val):
            continue
        try:
            giorno = datetime.strptime(str(riga[i_giorno]).strip()[:10], "%d/%m/%Y").date()
            ora = int(str(riga[i_da]).strip()[:2])
        except (TypeError, ValueError):
            scartate += 1
            continue
        valore = _numero(riga[i_val])
        if valore is None or not 0 <= ora <= 23:
            scartate += 1
            continue
        quarti[(giorno, ora)].append(valore)
        if i_k is not None and riga[i_k] is not None:
            costanti.setdefault(giorno, str(riga[i_k]).strip())

    if scartate:
        _LOGGER.warning("Excel mensile SET: %d righe non riconosciute scartate", scartate)

    per_giorno: dict[date, list[dict[str, Any]]] = defaultdict(list)
    for (giorno, ora), valori in sorted(quarti.items()):
        per_giorno[giorno].append(
            {
                "year": giorno.year,
                "month": giorno.month,
                "day": giorno.day,
                "hour": ora,
                "quarters": valori,
                "total": round(sum(valori), 6),
                "estimated": False,
                "kConstant": costanti.get(giorno, "1"),
            }
        )
    return dict(per_giorno)

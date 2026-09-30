"""Test del parser dell'export Excel mensile (distributors/set/excel.py).

Il file viene generato qui con openpyxl riproducendo ESATTAMENTE il layout
osservato il 30/09/2026 (intestazioni, celle tutte stringhe, virgola
decimale, foglio reattiva da ignorare), con valori di fantasia.
"""
from __future__ import annotations

import io
from datetime import date

import openpyxl
import pytest

from custom_components.contatore_letture.distributors.set import excel

POD = "IT221E00000000001"


def _xlsx(giorni: dict[date, list[float]], *, con_intestazione: bool = True,
          nome_foglio: str = "Energia Attiva Prevelata") -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = nome_foglio
    ws.append(["POD", "Ruolo", "Unita' di misura", "Anno", "Mese"])
    ws.append([POD, "Energia Attiva Prevelata", "kWh", "2026", "agosto"])
    ws.append([])
    if con_intestazione:
        ws.append(["Giorno", "Da ora", "A ora", "Valore", "Costante K", "Matricola Contatore"])
    for giorno, valori in sorted(giorni.items()):
        for i, v in enumerate(valori):
            da, a = i * 15, (i + 1) * 15
            ws.append([
                giorno.strftime("%d/%m/%Y"), f"{da // 60:02d}:{da % 60:02d}",
                f"{(a // 60) % 24:02d}:{a % 60:02d}", f"{v:.3f}".replace(".", ","),
                "1", "matricola-di-fantasia",
            ])
    ws2 = wb.create_sheet("Energia Reattiva Q1")
    ws2.append(["Giorno", "Da ora", "A ora", "Valore", "Costante K", "Matricola Contatore"])
    ws2.append(["01/08/2026", "00:00", "00:15", "9,999", "1", "matricola-di-fantasia"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_mese_completo_diventa_viste_giorno():
    g1, g2 = date(2026, 8, 1), date(2026, 8, 2)
    viste = excel.parse_excel_mensile(_xlsx({g1: [0.1] * 96, g2: [0.2] * 96}))
    assert set(viste) == {g1, g2}
    assert len(viste[g1]) == 24
    ora0 = viste[g1][0]
    assert (ora0["year"], ora0["month"], ora0["day"]) == (2026, 8, 1)
    assert ora0["hour"] == 0
    assert ora0["quarters"] == [0.1, 0.1, 0.1, 0.1]
    assert ora0["total"] == pytest.approx(0.4)
    assert ora0["estimated"] is False
    assert ora0["kConstant"] == "1"
    assert [e["hour"] for e in viste[g2]] == list(range(24))


def test_valori_con_virgola_e_ore_giuste():
    g = date(2026, 8, 20)
    valori = [0.0] * 96
    valori[13 * 4:13 * 4 + 4] = [0.117, 0.124, 0.289, 0.046]  # le 13:00-14:00
    viste = excel.parse_excel_mensile(_xlsx({g: valori}))
    ora13 = next(e for e in viste[g] if e["hour"] == 13)
    assert ora13["quarters"] == pytest.approx([0.117, 0.124, 0.289, 0.046])
    assert ora13["total"] == pytest.approx(0.576)


def test_il_foglio_reattiva_viene_ignorato():
    g = date(2026, 8, 1)
    viste = excel.parse_excel_mensile(_xlsx({g: [0.1] * 96}))
    assert sum(e["total"] for e in viste[g]) == pytest.approx(9.6)  # non 9.6 + 9.999


def test_compatibile_con_giorno_pubblicato_e_statistics():
    from custom_components.contatore_letture.distributors.set.api import giorno_pubblicato
    from custom_components.contatore_letture.distributors.set.statistics import kwh_del_giorno

    g = date(2026, 8, 1)
    viste = excel.parse_excel_mensile(_xlsx({g: [0.25] * 96}))
    assert giorno_pubblicato(viste[g])
    assert kwh_del_giorno(viste[g]) == pytest.approx(24.0)


def test_senza_intestazione_solleva_errore():
    with pytest.raises(excel.SetExcelError, match="intestazione"):
        excel.parse_excel_mensile(_xlsx({date(2026, 8, 1): [0.1] * 4}, con_intestazione=False))


def test_senza_foglio_attiva_solleva_errore():
    with pytest.raises(excel.SetExcelError, match="Foglio"):
        excel.parse_excel_mensile(_xlsx({date(2026, 8, 1): [0.1] * 4}, nome_foglio="Altro"))


def test_file_non_xlsx_solleva_errore():
    with pytest.raises(excel.SetExcelError):
        excel.parse_excel_mensile(b"non sono un xlsx")


def test_righe_non_riconosciute_vengono_scartate_senza_fallire(caplog):
    g = date(2026, 8, 1)
    wb = openpyxl.load_workbook(io.BytesIO(_xlsx({g: [0.1] * 8})))
    ws = wb["Energia Attiva Prevelata"]
    ws.append(["non-una-data", "xx", "yy", "abc", "1", "m"])
    buf = io.BytesIO()
    wb.save(buf)
    viste = excel.parse_excel_mensile(buf.getvalue())
    assert len(viste[g]) == 2
    assert "scartate" in caplog.text
